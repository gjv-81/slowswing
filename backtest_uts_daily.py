#!/usr/bin/env python3
"""
backtest_uts_daily.py — does the UTS daily swing model make money?
====================================================================
Backtests the UTS Aggressive daily BUY model over the full daily_data/
history (~2 years), and answers the open design question with data:

  ENTRY — tested HEAD TO HEAD, two definitions:
    continuation : enter on EVERY bar where the engine prints 'signal'
                   (core conditions met + breakout today). A trending
                   stock fires this over and over, all the way up.
    fresh_cross  : enter ONLY on the bar where core_ok FIRST turns true
                   (true today, false the prior bar) — the genuine
                   fresh qualification. One entry per qualifying run.

  EXIT — UTS long-exit conditions (uts_exit.py): the first bar after
         entry where cloud+vstop break, price-structure break, or
         momentum rollover fires.

  MECHANICS — no lookahead:
    - signal computed on a daily close (bar i)
    - ENTER at bar i+1 OPEN  (cannot trade the close you signal on)
    - EXIT  at the close of the bar where the UTS exit fires
    - ONE open position per ticker at a time (no pyramiding the
      same name); new signals while in a trade are ignored
    - trades still open at the last bar are closed there and marked
      'open_at_end' so they are not counted as clean exits

Run:
  cd ~/Documents/STS/15min && caffeinate -i python3.12 backtest_uts_daily.py
Output:
  uts_daily_backtest.xlsx   — per-trade log + summary, both entry modes
"""

import sys, warnings
warnings.filterwarnings('ignore')
from pathlib import Path
from datetime import datetime

import numpy as np
import pandas as pd

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
sys.path.insert(0, str(SCRIPT_DIR.parent))

from download_daily_data import load_daily_csv, OUT_DIR as DAILY_DIR
from uts_daily_engine import compute_uts_daily_buy, MIN_BARS_DAILY
from uts_exit import compute_uts_exit_series, uts_exit_bar

OUTPUT = str(SCRIPT_DIR / 'uts_daily_backtest.xlsx')

# entry definitions to test
ENTRY_MODES = ['continuation', 'fresh_cross']


# ─── ONE TICKER, ONE ENTRY MODE ──────────────────────────────────────────────

def backtest_one(ticker, df, states, exit_series, entry_mode):
    """Walk one ticker's history; return a list of trade dicts.

    states      — output of compute_uts_daily_buy(df)
    exit_series — output of compute_uts_exit_series(df)
    entry_mode  — 'continuation' or 'fresh_cross'
    """
    n = len(df)
    op = df['Open'].values
    cl = df['Close'].values
    hi = df['High'].values
    dates = df.index

    core = states['core_ok'].values
    sig = states['signal'].values

    # build the set of bars that count as an ENTRY TRIGGER for this mode
    if entry_mode == 'continuation':
        trigger = sig.copy()
    else:  # fresh_cross — core_ok true today, false (or absent) yesterday
        trigger = np.zeros(n, dtype=bool)
        for i in range(1, n):
            if core[i] and not core[i - 1]:
                trigger[i] = True

    trades = []
    i = 0
    while i < n:
        if not trigger[i]:
            i += 1
            continue

        # signal on bar i -> enter at bar i+1 open
        entry_pos = i + 1
        if entry_pos >= n:
            break  # signalled on the very last bar; no bar to enter on
        entry_price = op[entry_pos]
        if not np.isfinite(entry_price) or entry_price <= 0:
            i += 1
            continue

        # find the UTS exit after entry
        exit_pos = uts_exit_bar(df, entry_pos, exit_series)
        if exit_pos is None:
            # no exit within data — close at the last bar, flag it
            exit_pos = n - 1
            exit_price = cl[exit_pos]
            open_at_end = True
        else:
            exit_price = cl[exit_pos]
            open_at_end = False

        gain = exit_price - entry_price
        gain_pct = gain / entry_price * 100.0
        bars_held = exit_pos - entry_pos

        # max favourable / adverse excursion across the hold
        seg_hi = hi[entry_pos:exit_pos + 1]
        seg_lo = df['Low'].values[entry_pos:exit_pos + 1]
        mfe = (np.max(seg_hi) - entry_price) if len(seg_hi) else 0.0
        mae = (np.min(seg_lo) - entry_price) if len(seg_lo) else 0.0

        trades.append({
            'ticker':      ticker,
            'entry_mode':  entry_mode,
            'signal_date': dates[i].date(),
            'entry_date':  dates[entry_pos].date(),
            'exit_date':   dates[exit_pos].date(),
            'entry':       round(entry_price, 2),
            'exit':        round(exit_price, 2),
            'gain':        round(gain, 2),
            'gain_pct':    round(gain_pct, 3),
            'bars_held':   bars_held,
            'mfe':         round(mfe, 2),
            'mae':         round(mae, 2),
            'win':         int(gain > 0),
            'open_at_end': open_at_end,
        })

        # one position per ticker at a time: resume scanning AFTER the exit
        i = exit_pos + 1

    return trades


# ─── SUMMARY ─────────────────────────────────────────────────────────────────

def summarize(trades, label):
    """Return a one-row dict of summary stats for a list of trades."""
    if not trades:
        return {'set': label, 'trades': 0}

    # exclude still-open trades from P&L stats (they have no real exit)
    closed = [t for t in trades if not t['open_at_end']]
    n_open = len(trades) - len(closed)
    use = closed if closed else trades

    gains = np.array([t['gain'] for t in use])
    pcts = np.array([t['gain_pct'] for t in use])
    wins = np.array([t['win'] for t in use])
    held = np.array([t['bars_held'] for t in use])
    mfe = np.array([t['mfe'] for t in use])
    mae = np.array([t['mae'] for t in use])

    win_pcts = pcts[wins == 1]
    loss_pcts = pcts[wins == 0]
    avg_win = float(np.mean(win_pcts)) if len(win_pcts) else 0.0
    avg_loss = float(np.mean(loss_pcts)) if len(loss_pcts) else 0.0
    # expectancy per trade in % terms
    wr = float(np.mean(wins))
    expectancy = wr * avg_win + (1 - wr) * avg_loss
    # profit factor: gross win $ / gross loss $
    gross_win = float(np.sum(gains[gains > 0]))
    gross_loss = float(abs(np.sum(gains[gains < 0])))
    pf = (gross_win / gross_loss) if gross_loss > 0 else float('inf')

    return {
        'set':            label,
        'trades':         len(use),
        'still_open':     n_open,
        'win%':           round(wr * 100, 1),
        'avg_gain%':      round(float(np.mean(pcts)), 3),
        'median_gain%':   round(float(np.median(pcts)), 3),
        'avg_win%':       round(avg_win, 3),
        'avg_loss%':      round(avg_loss, 3),
        'expectancy%':    round(expectancy, 3),
        'profit_factor':  round(pf, 2) if np.isfinite(pf) else None,
        'total_$':        round(float(np.sum(gains)), 2),
        'avg_bars_held':  round(float(np.mean(held)), 1),
        'avg_mfe$':       round(float(np.mean(mfe)), 2),
        'avg_mae$':       round(float(np.mean(mae)), 2),
    }


# ─── EXCEL OUTPUT ────────────────────────────────────────────────────────────

def write_excel(all_trades, summaries, per_ticker, coverage, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    C_HEAD = '1F3864'
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(name, df):
        if df is None or df.empty:
            ws = wb.create_sheet(name)
            ws['A1'] = '(no rows)'
            return
        ws = wb.create_sheet(name)
        ws.freeze_panes = 'A2'
        for c, col in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=c, value=str(col))
            cell.font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
            cell.fill = PatternFill('solid', start_color=C_HEAD)
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = border
        ws.row_dimensions[1].height = 28
        for r, (_, row) in enumerate(df.iterrows(), 2):
            for c, val in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.font = Font(name='Arial', size=9)
                cell.alignment = Alignment(horizontal='center')
                cell.border = border
        for c, col in enumerate(df.columns, 1):
            w = max(len(str(col)),
                    max((len(str(df.iloc[r, c - 1]))
                         for r in range(len(df))), default=0))
            ws.column_dimensions[get_column_letter(c)].width = min(w + 3, 24)

    # README
    ws0 = wb.create_sheet('README', 0)
    ws0['A1'] = 'UTS Daily Swing Backtest'
    ws0['A1'].font = Font(bold=True, size=15, color=C_HEAD, name='Arial')
    info = [
        ('Generated', datetime.now().strftime('%Y-%m-%d %H:%M')),
        ('Model', 'UTS Aggressive daily BUY -> UTS exit'),
        ('Entry', 'enter at bar i+1 OPEN after a signal on bar i'),
        ('Exit', 'UTS long-exit conditions; exit at that bar CLOSE'),
        ('Rule', 'one open position per ticker at a time'),
        ('continuation', 'enter on EVERY signal bar (fires up a whole trend)'),
        ('fresh_cross', 'enter only on the first qualifying bar of a run'),
        ('Note', 'still-open trades are excluded from P&L stats'),
        ('Tickers tested', coverage.get('tickers_used', 0)),
        ('Data range', coverage.get('range', 'n/a')),
    ]
    for r, (k, v) in enumerate(info, 3):
        ws0[f'A{r}'] = k
        ws0[f'A{r}'].font = Font(bold=True, name='Arial', size=10)
        ws0[f'B{r}'] = str(v)
        ws0[f'B{r}'].font = Font(name='Arial', size=10)
    ws0.column_dimensions['A'].width = 18
    ws0.column_dimensions['B'].width = 70

    sheet('Summary', pd.DataFrame(summaries))
    sheet('Per-Ticker', per_ticker)
    for mode in ENTRY_MODES:
        rows = [t for t in all_trades if t['entry_mode'] == mode]
        sheet(f'Trades {mode}', pd.DataFrame(rows))

    wb.save(path)


# ─── MAIN ────────────────────────────────────────────────────────────────────

def main():
    print('=' * 64)
    print('  UTS DAILY SWING BACKTEST')
    print('  continuation vs fresh-cross entry  |  UTS exit')
    print('=' * 64)

    csvs = sorted(DAILY_DIR.glob('*.csv'))
    if not csvs:
        print(f'  No CSVs in {DAILY_DIR} — run download_daily_data.py first.')
        sys.exit(1)
    tickers = [p.stem for p in csvs]
    print(f'  {len(tickers)} tickers in {DAILY_DIR}\n')

    all_trades = []
    per_ticker_rows = []
    used = 0
    skipped = 0
    dmin, dmax = None, None

    for idx, ticker in enumerate(tickers, 1):
        df = load_daily_csv(ticker)
        if df is None or len(df) < MIN_BARS_DAILY:
            skipped += 1
            continue

        # track overall data range
        if dmin is None or df.index[0] < dmin:
            dmin = df.index[0]
        if dmax is None or df.index[-1] > dmax:
            dmax = df.index[-1]

        # compute engine + exit ONCE per ticker, reuse for both modes
        states = compute_uts_daily_buy(df)
        exit_series = compute_uts_exit_series(df)

        tk_trades = {}
        for mode in ENTRY_MODES:
            trades = backtest_one(ticker, df, states, exit_series, mode)
            all_trades.extend(trades)
            tk_trades[mode] = trades

        # per-ticker row: P&L for each mode side by side
        row = {'ticker': ticker, 'bars': len(df)}
        for mode in ENTRY_MODES:
            closed = [t for t in tk_trades[mode] if not t['open_at_end']]
            row[f'{mode}_trades'] = len(closed)
            if closed:
                g = [t['gain'] for t in closed]
                w = [t['win'] for t in closed]
                row[f'{mode}_win%'] = round(np.mean(w) * 100, 1)
                row[f'{mode}_total$'] = round(float(np.sum(g)), 2)
            else:
                row[f'{mode}_win%'] = None
                row[f'{mode}_total$'] = None
        per_ticker_rows.append(row)
        used += 1

        if idx % 100 == 0:
            print(f'  {idx}/{len(tickers)} processed, '
                  f'{len(all_trades)} trades so far')

    print(f'\n  Done: {used} tickers tested, {skipped} skipped '
          f'(insufficient bars)')

    coverage = {
        'tickers_used': used,
        'range': (f'{dmin.date()} -> {dmax.date()}'
                  if dmin is not None else 'n/a'),
    }

    # summaries — overall per mode
    summaries = []
    for mode in ENTRY_MODES:
        rows = [t for t in all_trades if t['entry_mode'] == mode]
        summaries.append(summarize(rows, mode))

    # print the head-to-head
    print('\n' + '=' * 64)
    print('  HEAD TO HEAD — continuation vs fresh-cross')
    print('=' * 64)
    keys = ['trades', 'still_open', 'win%', 'avg_gain%', 'median_gain%',
            'expectancy%', 'profit_factor', 'total_$', 'avg_bars_held']
    print(f'  {"metric":<18}{"continuation":>16}{"fresh_cross":>16}')
    print('  ' + '-' * 48)
    s_by_mode = {s['set']: s for s in summaries}
    for k in keys:
        cv = s_by_mode['continuation'].get(k, '-')
        fc = s_by_mode['fresh_cross'].get(k, '-')
        print(f'  {k:<18}{str(cv):>16}{str(fc):>16}')
    print('=' * 64)

    per_ticker = pd.DataFrame(per_ticker_rows)
    write_excel(all_trades, summaries, per_ticker, coverage, OUTPUT)
    print(f'\n  Saved: {OUTPUT}')
    print('  Sheets: README, Summary, Per-Ticker, Trades (per mode)')
    print('\n  NOTE: this is a backtest on ~2yr of data. It tells you')
    print('  whether the model HAS made money, not whether it WILL.')
    print('  Read the win%, expectancy and profit_factor before')
    print('  trusting any single name the live scanner flags.')


# ─── SELF-TEST ───────────────────────────────────────────────────────────────

def _selftest():
    print('  backtest_uts_daily — self-test\n' + '=' * 50)
    import pytz
    ET = pytz.timezone('US/Eastern')
    rng = np.random.RandomState(7)
    checks = []

    # synthetic ticker: base -> uptrend -> rollover
    p = [100.0]
    for i in range(1, 400):
        if i < 150:
            d = rng.normal(0.0, 0.004)
        elif i < 300:
            d = rng.normal(0.011, 0.005)
        else:
            d = rng.normal(-0.006, 0.005)
        p.append(max(1.0, p[-1] * (1 + d)))
    p = np.array(p)
    df = pd.DataFrame({
        'Open':  p * (1 + rng.normal(0, 0.001, 400)),
        'High':  p * (1 + np.abs(rng.normal(0, 0.004, 400))),
        'Low':   p * (1 - np.abs(rng.normal(0, 0.004, 400))),
        'Close': p,
        'Volume': rng.randint(1e6, 5e6, 400),
    }, index=pd.date_range('2023-01-02', periods=400, freq='B', tz=ET))

    states = compute_uts_daily_buy(df)
    exit_series = compute_uts_exit_series(df)

    cont = backtest_one('TEST', df, states, exit_series, 'continuation')
    fresh = backtest_one('TEST', df, states, exit_series, 'fresh_cross')

    checks.append(('continuation produces trades', len(cont) > 0))
    checks.append(('fresh_cross produces trades', len(fresh) > 0))
    checks.append(('fresh_cross has <= continuation trade count',
                   len(fresh) <= len(cont)))
    # no lookahead: entry_date strictly after signal_date for every trade
    ok_dates = all(t['entry_date'] > t['signal_date']
                   for t in cont + fresh)
    checks.append(('entry always after signal (no lookahead)', ok_dates))
    # exit not before entry
    ok_exit = all(t['exit_date'] >= t['entry_date'] for t in cont + fresh)
    checks.append(('exit never before entry', ok_exit))
    # one position at a time: trades within a mode never overlap in bars
    def no_overlap(trades):
        s = sorted(trades, key=lambda t: t['entry_date'])
        for a, b in zip(s, s[1:]):
            if b['entry_date'] < a['exit_date']:
                return False
        return True
    checks.append(('continuation: positions never overlap',
                   no_overlap(cont)))
    checks.append(('fresh_cross: positions never overlap',
                   no_overlap(fresh)))
    # summarize runs without error and excludes open trades
    s = summarize(cont, 'continuation')
    checks.append(('summarize returns a dict with trades key',
                   isinstance(s, dict) and 'trades' in s))
    # win flag consistent with gain sign
    checks.append(('win flag matches gain sign',
                   all(t['win'] == int(t['gain'] > 0)
                       for t in cont + fresh)))
    # gain math: exit - entry
    checks.append(('gain == exit - entry (to cents)',
                   all(abs(t['gain'] - (t['exit'] - t['entry'])) < 0.011
                       for t in cont + fresh)))
    # downtrend-only frame: should yield few/no trades, never crash
    dn = [100.0]
    for i in range(1, 300):
        dn.append(max(1.0, dn[-1] * (1 + rng.normal(-0.005, 0.004))))
    dn = np.array(dn)
    df_dn = pd.DataFrame({
        'Open': dn, 'High': dn * 1.003, 'Low': dn * 0.997,
        'Close': dn, 'Volume': rng.randint(1e6, 5e6, 300),
    }, index=pd.date_range('2023-01-02', periods=300, freq='B', tz=ET))
    st_dn = compute_uts_daily_buy(df_dn)
    ex_dn = compute_uts_exit_series(df_dn)
    dn_trades = backtest_one('DN', df_dn, st_dn, ex_dn, 'continuation')
    checks.append(('downtrend frame does not crash',
                   isinstance(dn_trades, list)))

    print()
    for name, ok in checks:
        print(f'  [{"PASS" if ok else "FAIL"}] {name}')
    n_pass = sum(1 for _, ok in checks if ok)
    print(f'\n  {n_pass}/{len(checks)} checks passed')
    if n_pass == len(checks):
        print(f'\n  (synthetic: continuation {len(cont)} trades, '
              f'fresh_cross {len(fresh)} trades)')
        print('  ' + '-' * 48)
        for s in (summarize(cont, 'continuation'),
                  summarize(fresh, 'fresh_cross')):
            print(f"  {s['set']:<14} win {s.get('win%')}%  "
                  f"exp {s.get('expectancy%')}%  "
                  f"PF {s.get('profit_factor')}  "
                  f"${s.get('total_$')}")
    return n_pass == len(checks)


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == 'selftest':
        ok = _selftest()
        sys.exit(0 if ok else 1)
    main()
